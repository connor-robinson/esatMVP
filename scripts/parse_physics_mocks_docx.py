# -*- coding: utf-8 -*-
"""Parse ESAT Physics Mock Modules DOCX into structured JSON."""
from __future__ import annotations

import json
import os
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX = Path(r"c:\Users\anson\Downloads\ESAT_Physics_Mock_Modules_A_B.docx")
OUT_DIR = Path(r"c:\Users\anson\Desktop\nocalcMVP2_real\tmp_physics_mocks")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def cell_text(cell) -> str:
    return cell.text.replace("\r\n", "\n").replace("\r", "\n").strip()


def dump_tables(doc: Document) -> None:
    out = OUT_DIR / "tables_full.txt"
    with out.open("w", encoding="utf-8") as f:
        for ti, table in enumerate(doc.tables):
            f.write(f"\n======== TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} ========\n")
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    text = cell_text(cell)
                    if text:
                        f.write(f"[T{ti}R{ri}C{ci}]\n{text}\n---\n")
    print("wrote", out)


def extract_images(doc_path: Path) -> None:
    img_dir = OUT_DIR / "images"
    img_dir.mkdir(exist_ok=True)
    with zipfile.ZipFile(doc_path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        print("media count", len(media))
        for n in media:
            data = z.read(n)
            dest = img_dir / Path(n).name
            dest.write_bytes(data)
            print(dest.name, len(data))


def iter_body_blocks(doc: Document):
    """Yield ('p', para) or ('t', table) in document order."""
    body = doc.element.body
    p_map = {p._element: p for p in doc.paragraphs}
    t_map = {t._element: t for t in doc.tables}
    for child in body.iterchildren():
        if child.tag == qn("w:p") and child in p_map:
            yield ("p", p_map[child])
        elif child.tag == qn("w:tbl") and child in t_map:
            yield ("t", t_map[child])


OPTION_RE = re.compile(r"^([A-H])\.\s*(.*)$", re.S)
Q_START_RE = re.compile(r"^(\d+)\.\s*(.*)$", re.S)


def parse_candidate_question(table) -> dict | None:
    """Parse a candidate question table (stem + options, possibly nested)."""
    # Flatten unique cell texts in reading order of first column-ish
    texts = []
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            t = cell_text(cell)
            if not t or t in seen:
                continue
            seen.add(t)
            texts.append(t)

    if not texts:
        return None

    joined = "\n".join(texts)
    # Some tables put everything in one cell with newlines
    lines = []
    for t in texts:
        lines.extend(t.split("\n"))

    # Find question number and stem
    first = texts[0]
    m = Q_START_RE.match(first)
    if not m:
        # maybe "1." alone then stem
        if re.match(r"^\d+\.?$", first.strip()) and len(texts) > 1:
            qnum = int(re.match(r"^(\d+)", first).group(1))
            stem_parts = []
            options = {}
            mode = "stem"
            for t in texts[1:]:
                om = OPTION_RE.match(t.strip())
                if om:
                    mode = "opts"
                    options[om.group(1)] = om.group(2).strip()
                elif mode == "stem":
                    stem_parts.append(t)
                else:
                    # continuation of last option?
                    last = list(options.keys())[-1] if options else None
                    if last:
                        options[last] = (options[last] + " " + t).strip()
            return {
                "number": qnum,
                "stem": "\n".join(stem_parts).strip(),
                "options": options,
                "raw_texts": texts,
            }
        return None

    qnum = int(m.group(1))
    rest = m.group(2).strip()
    options = {}
    stem_lines = []
    if rest:
        stem_lines.append(rest)

    # Process remaining lines/texts for options
    all_lines = []
    for t in texts:
        for line in t.split("\n"):
            all_lines.append(line)

    # Rebuild from all_lines more carefully
    stem_parts = []
    options = {}
    started = False
    mode = "stem"
    for line in all_lines:
        s = line.strip()
        if not started:
            qm = Q_START_RE.match(s)
            if qm:
                started = True
                rem = qm.group(2).strip()
                if rem:
                    stem_parts.append(rem)
                continue
            continue
        om = OPTION_RE.match(s)
        if om:
            mode = "opts"
            options[om.group(1)] = om.group(2).strip()
        elif mode == "stem":
            if s:
                stem_parts.append(s)
        elif mode == "opts" and s and options:
            last = list(options.keys())[-1]
            options[last] = (options[last] + "\n" + s).strip()

    return {
        "number": qnum,
        "stem": "\n".join(stem_parts).strip(),
        "options": options,
        "raw_texts": texts,
    }


def parse_editor_key_block(texts: list[str]) -> dict | None:
    """Parse editor key solution block from table cell texts."""
    blob = "\n".join(texts)
    # Typical header: A1  ★ EDITOR PICK or A1
    header = texts[0] if texts else ""
    hm = re.match(r"^([AB])(\d+)\s*(.*)$", header.strip())
    if not hm:
        # sometimes first line is like "A1 ★ EDITOR PICK"
        hm = re.search(r"\b([AB])(\d+)\b", blob[:80])
        if not hm:
            return None
        module = hm.group(1)
        number = int(hm.group(2))
        editor_pick = "EDITOR PICK" in blob[:120].upper() or "★" in blob[:120]
    else:
        module = hm.group(1)
        number = int(hm.group(2))
        editor_pick = "EDITOR PICK" in hm.group(3).upper() or "★" in hm.group(3)

    def extract_field(label: str) -> str | None:
        # Match Label: rest until next known label
        labels = [
            "Answer",
            "Topic",
            "Difficulty",
            "Target",
            "Tip",
            "Solution",
            "Distractor map",
            "Benchmark",
        ]
        # Allow variations
        pat = rf"(?:^|\n)\s*{re.escape(label)}\s*[:：]\s*"
        m = re.search(pat, blob, re.I)
        if not m:
            return None
        start = m.end()
        # find next label
        next_positions = []
        for lab in labels:
            if lab.lower() == label.lower():
                continue
            nm = re.search(rf"(?:^|\n)\s*{re.escape(lab)}\s*[:：]", blob[start:], re.I)
            if nm:
                next_positions.append(start + nm.start())
        end = min(next_positions) if next_positions else len(blob)
        return blob[start:end].strip()

    answer = extract_field("Answer")
    topic = extract_field("Topic")
    difficulty = extract_field("Difficulty")
    target = extract_field("Target") or extract_field("Target time")
    tip = extract_field("Tip")
    solution = extract_field("Solution")
    distractor_raw = extract_field("Distractor map")
    benchmark = extract_field("Benchmark") or extract_field("Benchmark note")

    distractors = {}
    if distractor_raw:
        # Lines like A — text or A: text or A. text
        for line in distractor_raw.split("\n"):
            line = line.strip()
            if not line:
                continue
            dm = re.match(r"^([A-H])\s*[—–\-:\.]\s*(.*)$", line)
            if dm:
                distractors[dm.group(1)] = dm.group(2).strip()
            elif distractors:
                last = list(distractors.keys())[-1]
                distractors[last] = (distractors[last] + " " + line).strip()

    return {
        "module": module,
        "number": number,
        "editorPick": editor_pick,
        "answer": answer,
        "topic": topic,
        "difficulty": difficulty,
        "target": target,
        "tip": tip,
        "solution": solution,
        "distractors": distractors,
        "distractorRaw": distractor_raw,
        "benchmarkNote": benchmark,
        "raw": texts,
    }


def main() -> None:
    doc = Document(str(DOCX))
    dump_tables(doc)
    extract_images(DOCX)

    # Walk document in order; split into candidate A, candidate B, editor A, editor B
    section = "pre"
    candidate_a = []
    candidate_b = []
    editor_a = []
    editor_b = []
    overview_tables = []

    for kind, node in iter_body_blocks(doc):
        if kind == "p":
            t = node.text.strip()
            if t == "PHYSICS MODULE A" or (t.startswith("PHYSICS MODULE A") and "Candidate" not in t):
                # next candidate
                pass
            if t == "Candidate Paper" and section in ("pre", "end_a", "mod_a_head"):
                section = "cand_a" if "cand_a" not in section or section.startswith("pre") or section == "mod_a_head" else section
            if "PHYSICS MODULE A" in t and "Candidate" not in t:
                section = "mod_a_head"
            if t == "Candidate Paper" and section == "mod_a_head":
                section = "cand_a"
            if t == "End of Module A":
                section = "end_a"
            if "PHYSICS MODULE B" in t and "Candidate" not in t:
                section = "mod_b_head"
            if t == "Candidate Paper" and section == "mod_b_head":
                section = "cand_b"
            if t == "End of Module B":
                section = "end_b"
            if t == "EDITOR KEY" or t.startswith("EDITOR KEY"):
                section = "editor"
            if t.startswith("Module A: worked solutions"):
                section = "editor_a"
            if t.startswith("Module B: worked solutions"):
                section = "editor_b"
            if t.startswith("Module A: answer overview"):
                section = "overview_a"
            if t.startswith("Module B: answer overview"):
                section = "overview_b"
            continue

        # table
        if section == "cand_a":
            q = parse_candidate_question(node)
            if q and q.get("number"):
                candidate_a.append(q)
        elif section == "cand_b":
            q = parse_candidate_question(node)
            if q and q.get("number"):
                candidate_b.append(q)
        elif section in ("editor_a", "overview_a"):
            texts = []
            seen = set()
            for row in node.rows:
                for cell in row.cells:
                    ct = cell_text(cell)
                    if ct and ct not in seen:
                        seen.add(ct)
                        texts.append(ct)
            if section == "overview_a":
                overview_tables.append(("A", texts))
            else:
                ek = parse_editor_key_block(texts)
                if ek:
                    editor_a.append(ek)
        elif section in ("editor_b", "overview_b"):
            texts = []
            seen = set()
            for row in node.rows:
                for cell in row.cells:
                    ct = cell_text(cell)
                    if ct and ct not in seen:
                        seen.add(ct)
                        texts.append(ct)
            if section == "overview_b":
                overview_tables.append(("B", texts))
            else:
                ek = parse_editor_key_block(texts)
                if ek:
                    editor_b.append(ek)

    result = {
        "candidateA": candidate_a,
        "candidateB": candidate_b,
        "editorA": editor_a,
        "editorB": editor_b,
        "overview": overview_tables,
        "counts": {
            "candA": len(candidate_a),
            "candB": len(candidate_b),
            "edA": len(editor_a),
            "edB": len(editor_b),
        },
    }
    (OUT_DIR / "parsed.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("counts", result["counts"])
    # Sanity print first question
    if candidate_a:
        print("A1 stem:", candidate_a[0]["stem"][:100])
        print("A1 opts:", list(candidate_a[0]["options"].keys()), list(candidate_a[0]["options"].values())[:2])
    if editor_a:
        print("EK A1:", {k: editor_a[0].get(k) for k in ("answer", "topic", "difficulty", "target", "editorPick")})
        print("distractors", editor_a[0].get("distractors"))


if __name__ == "__main__":
    main()
